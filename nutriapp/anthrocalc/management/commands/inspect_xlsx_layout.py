"""Work out what shape a workbook is, and draft a profile for loading it.

    ./manage.py inspect_xlsx_layout --file book.xlsx --sheet 'Hoja 1'
    ./manage.py inspect_xlsx_layout --file book.xlsx --sheet 'Hoja 1' --write mylayout.json
    ./manage.py inspect_xlsx_layout --dump nimacabaj
    ./manage.py inspect_xlsx_layout --file book.xlsx --compare nimacabaj

The draft is a starting point, never a finished profile: read the notes, fix
the file, then load with ``load_xlsx_patients --profile mylayout.json``.
``--compare`` measures the detector against a profile that is known to be
right, which is how the accuracy figures in
``docs/features/bulk_patients/profiles.md`` were arrived at.
"""

from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError

from anthrocalc import bulk_import, bulk_layout

LEVEL_ORDER = ("missing", "check", "found")


class Command(BaseCommand):
    help = "Detect the layout of a workbook and draft a loader profile for it."

    def add_arguments(self, parser):
        parser.add_argument("--file", help="Path to the .xlsx workbook.")
        parser.add_argument(
            "--sheet",
            default=None,
            help="Worksheet to inspect. Defaults to every sheet in the workbook.",
        )
        parser.add_argument(
            "--write", default=None, help="Write the drafted profile to this JSON path."
        )
        parser.add_argument(
            "--dump",
            default=None,
            help="Print a built-in profile as JSON instead of detecting anything.",
        )
        parser.add_argument(
            "--compare",
            default=None,
            help="Built-in profile (or JSON path) to check the detected layout against.",
        )
        parser.add_argument(
            "--name", default="detected", help="Name to give the drafted profile."
        )

    def handle(self, *args, **options):
        if options["dump"]:
            try:
                self.stdout.write(bulk_import.load_profile(options["dump"]).to_json())
            except bulk_import.ProfileError as exc:
                raise CommandError(str(exc)) from exc
            return

        if not options["file"]:
            raise CommandError("give me a --file to inspect, or a --dump profile name")

        try:
            workbook = openpyxl.load_workbook(options["file"], data_only=True)
        except OSError as exc:
            raise CommandError(str(exc)) from exc

        sheets = [options["sheet"]] if options["sheet"] else workbook.sheetnames
        for sheet_name in sheets:
            if sheet_name not in workbook.sheetnames:
                raise CommandError(f"sheet {sheet_name!r} not in {workbook.sheetnames}")
            sheet = workbook[sheet_name]
            # A workbook can carry chart-only sheets, which have no cells.
            if not hasattr(sheet, "cell"):
                self.stdout.write(f"\n{sheet_name!r} — a chart sheet, nothing to read")
                continue
            self.__inspect(sheet, options)

    def __inspect(self, sheet, options):
        self.stdout.write(
            self.style.MIGRATE_HEADING(
                f"\n{sheet.title!r} — {sheet.max_row} rows × {sheet.max_column} columns"
            )
        )
        detection = bulk_layout.detect(sheet, profile_name=options["name"])
        profile = detection.profile

        for level in LEVEL_ORDER:
            for note in detection.of_level(level):
                style = {
                    "missing": self.style.ERROR,
                    "check": self.style.WARNING,
                    "found": lambda text: text,
                }[level]
                self.stdout.write("  " + style(str(note)))

        self.stdout.write(
            f"\n  {len(profile.blocks)} rounds and "
            f"{sum(len(g.rows) for g in profile.row_groups)} candidate rows detected."
        )
        for block in profile.blocks:
            source = (
                f"date column {block.date_col}"
                if block.date_col
                else f"round date at row {block.round_date_row}, column {block.round_date_col}"
                if block.round_date_col
                else self.style.ERROR("no date")
            )
            self.stdout.write(
                f"    round {block.number:>4}  weight {block.weight_col:>3}  "
                f"height {block.height_col:>3}  {source}"
            )

        problems = profile.problems()
        if problems:
            self.stdout.write(self.style.ERROR("\n  The draft is not loadable as it stands:"))
            for problem in problems:
                self.stdout.write(self.style.ERROR(f"    {problem}"))

        if options["compare"]:
            self.__compare(profile, options["compare"])

        if options["write"]:
            path = Path(options["write"])
            path.write_text(profile.to_json(), encoding="utf-8")
            self.stdout.write(
                self.style.SUCCESS(f"\n  Draft written to {path}. Read it before loading with it.")
            )

    def __compare(self, detected, reference):
        try:
            expected = bulk_import.load_profile(reference)
        except bulk_import.ProfileError as exc:
            raise CommandError(str(exc)) from exc

        differences = bulk_layout.compare(detected, expected)
        self.stdout.write(
            self.style.MIGRATE_HEADING(f"\n  Against the {expected.name!r} profile")
        )
        if not differences:
            self.stdout.write(self.style.SUCCESS("    identical"))
            return
        for difference in differences:
            self.stdout.write(self.style.WARNING(f"    {difference}"))
        self.stdout.write(f"    {len(differences)} differences.")
